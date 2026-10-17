# Copyright 2026 Michael Brackx
# SPDX-License-Identifier: Apache-2.0
"""Native DOCX specimens with complete, editable embedded fonts.

Requires python-docx in addition to the shared builder's FontTools dependency.
Font embedding follows ECMA-376 Part 1 section 17.8.1; OpenType settings use
the Office 2010 WordprocessingML extensions. Word rendering needs separate QA.
"""

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from math import ceil
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Cm, Pt, RGBColor
from fontTools.ttLib import TTFont
from lxml import etree


FONT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.obfuscatedFont"


def obfuscate_font(data, key):
    """XOR the first 32 bytes with the reversed GUID bytes, repeated twice."""
    mask = UUID(key).bytes[::-1]
    result = bytearray(data)
    for index in range(32):
        result[index] ^= mask[index % 16]
    return bytes(result)


def set_value(parent, tag, value):
    node = OxmlElement(tag)
    node.set(qn(tag.split(":")[0] + ":val"), str(value))
    parent.append(node)
    return node


def set_font(properties, family):
    fonts = properties.get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), family)
    # Theme font attributes take precedence over the explicit font names.
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + name), None)


def embed_fonts(document, paths, family_for):
    part = document.part.part_related_by(RT.FONT_TABLE)
    table = OxmlElement("w:fonts")
    for index, path in enumerate(paths, 1):
        family = family_for(path)  # Also checks editable embedding rights.
        data = path.read_bytes()
        key = "{" + str(UUID(bytes=sha256(data).digest()[:16], version=5)).upper() + "}"
        font_part = Part(PackURI(f"/word/fonts/font{index}.odttf"),
                         FONT_CONTENT_TYPE, obfuscate_font(data, key), part.package)
        relation = part.relate_to(font_part, RT.FONT)
        entry = OxmlElement("w:font")
        entry.set(qn("w:name"), family)
        with TTFont(path) as font:
            os2 = font["OS/2"]
            set_value(entry, "w:charset", "00")
            set_value(entry, "w:family", "auto")
            set_value(entry, "w:pitch", "variable")
            signature = OxmlElement("w:sig")
            for attr, field in (("usb0", "ulUnicodeRange1"), ("usb1", "ulUnicodeRange2"),
                                ("usb2", "ulUnicodeRange3"), ("usb3", "ulUnicodeRange4"),
                                ("csb0", "ulCodePageRange1"), ("csb1", "ulCodePageRange2")):
                signature.set(qn("w:" + attr), f"{getattr(os2, field, 0):08X}")
            entry.append(signature)
        embedded = OxmlElement("w:embedRegular")
        embedded.set(qn("r:id"), relation)
        embedded.set(qn("w:fontKey"), key)
        embedded.set(qn("w:subsetted"), "false")
        entry.append(embedded)
        table.append(entry)
    part._blob = etree.tostring(table, xml_declaration=True, encoding="UTF-8", standalone=True)
    settings = document.settings.element
    # Keep all glyphs and all GSUB lookups available when a recipient edits.
    anchor = settings.find(qn("w:proofState"))
    if anchor is None:
        anchor = settings.find(qn("w:defaultTabStop"))
    for name, value in (("embedTrueTypeFonts", "true"), ("embedSystemFonts", "true"),
                        ("saveSubsetFonts", "false")):
        node = OxmlElement("w:" + name)
        node.set(qn("w:val"), value)
        if anchor is not None:
            anchor.addprevious(node)
        else:
            settings.append(node)
    for setting in settings.findall("w:compat/w:compatSetting", settings.nsmap):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")


def build(spec, root, base, family_for, output=None):
    font_path = root / spec.folder / spec.font
    family = family_for(font_path)
    body_family = family_for(base)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = body_family, Pt(10)
    normal.font.color.rgb = RGBColor.from_string("24333F")
    set_font(normal.element.get_or_add_rPr(), body_family)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in (("Title", 24), ("Heading 1", 12)):
        style = document.styles[name]
        style.font.name, style.font.size = body_family, Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style.element.get_or_add_rPr()
        set_font(properties, body_family)
        properties.color.attrib.pop(qn("w:themeColor"), None)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            borders.getparent().remove(borders)

    for name, size in (("Source", 9), ("Rendered", spec.size), ("Caption", 9)):
        style = (document.styles[name] if name in document.styles else
                 document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH))
        style.base_style = normal
        style.font.name = family if name == "Rendered" else body_family
        style.font.size = Pt(size)
        style.font.bold = style.font.italic = False
        style.font.color.rgb = RGBColor.from_string("5E707C" if name == "Caption" else "24333F")
        set_font(style.element.get_or_add_rPr(), style.font.name)
        style.element.rPr.color.attrib.pop(qn("w:themeColor"), None)
        style.paragraph_format.space_after = Pt(6)
    rendered = document.styles["Rendered"]
    rendered.paragraph_format.line_spacing = 1.0
    rendered.paragraph_format.keep_together = True
    # 'atLeast' lets tall glyph compositions grow rather than clipping them.
    if isinstance(spec.line_height, int):
        from docx.enum.text import WD_LINE_SPACING
        rendered.paragraph_format.line_spacing = Pt(spec.line_height)
        rendered.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST

    def paragraph(text, style=None):
        return document.add_paragraph(text, style)

    paragraph(spec.title, "Title")
    paragraph(spec.description)
    paragraph("Each pair contains the same editable input in two embedded fonts. "
              "Edit the computational example to change its result; the source copy is independent.")
    for example in spec.examples:
        paragraph(example.label, "Heading 1")
        if not spec.source_appendix:
            paragraph("Underlying text in Roboto", "Caption")
            paragraph(example.source, "Source")
        paragraph("Editable computational font", "Caption")
        if spec.unbroken_run:
            # Match the ODT's wide shaping paragraph inside a page-width canvas.
            # A layout engine may otherwise break input before GSUB collapses it.
            # Word may normalize extreme indents; these examples need Word QA.
            table = document.add_table(rows=1, cols=1)
            table.autofit = False
            table.columns[0].width = Cm(17)
            cell = table.cell(0, 0)
            cell_properties = cell._tc.get_or_add_tcPr()
            width = cell_properties.find(qn("w:tcW"))
            width.set(qn("w:type"), "dxa")
            width.set(qn("w:w"), str(Cm(17).twips))
            p = cell.paragraphs[0]
            p.style = "Rendered"
            p.add_run(example.source)
            shaping_margin = max(2000, ceil(len(example.source) * spec.size * 2.54 / 72) + 2000)
            # Write twips directly: python-docx validates the intermediate EMU
            # length as int32, although OOXML stores this value as signed twips.
            p._p.get_or_add_pPr().get_or_add_ind().set(qn("w:right"), str(Cm(-shaping_margin).twips))
        else:
            p = paragraph(example.source, "Rendered")
        p.paragraph_format.keep_with_next = True
        ppr = p._p.get_or_add_pPr()
        suppress = OxmlElement("w:suppressAutoHyphens")
        suppress.set(qn("w:val"), "true")
        ppr.insert_element_before(suppress, "w:spacing", "w:ind")
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "EAF3F8")
        if spec.unbroken_run:
            cell_properties.append(shade)
        else:
            ppr.insert_element_before(shade, "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind")
        run = p.runs[0]
        properties = run._r.get_or_add_rPr()
        set_font(properties, family)
        properties.get_or_add_noProof().val = True
        set_value(properties, "w14:ligatures", "standardContextual")
        set_value(properties, "w14:cntxtAlts", "true")
        paragraph("Expected: " + example.expected, "Caption")
    paragraph("Try it", "Heading 1")
    paragraph(spec.exercise)
    paragraph(spec.note, "Caption")
    paragraph("Use desktop Word with standard ligatures and contextual alternates enabled. "
              "Keep each input in one paragraph with consistent formatting. "
              "Copying the result copies the original input. Fonts are embedded in full.", "Caption")
    if spec.source_appendix:
        for example in spec.examples:
            p = paragraph(example.label + " plain text", "Heading 1")
            p.paragraph_format.page_break_before = True
            paragraph(example.source, "Source")

    embed_fonts(document, (base, font_path), family_for)
    for filename, target in (("LICENSE", "LICENSE.txt"),
                             ("THIRD_PARTY_NOTICES.md", "THIRD_PARTY_NOTICES.txt")):
        part = Part(PackURI("/" + target), "text/plain",
                    (root / filename).read_bytes(), document.part.package)
        document.part.relate_to(part, "urn:computational-fonts:relationships:license")
    document.core_properties.title = spec.title
    document.core_properties.author = "computational-fonts"
    document.core_properties.last_modified_by = "computational-fonts"
    fixed_date = datetime(2026, 1, 1, tzinfo=timezone.utc)
    document.core_properties.created = document.core_properties.modified = fixed_date
    buffer = BytesIO()
    document.save(buffer)
    output = output or font_path.with_name(font_path.stem.removesuffix("-font") + ".docx")
    # python-docx uses wall-clock ZIP timestamps; normalize for reproducibility.
    with ZipFile(buffer) as source, ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            target.writestr(info, source.read(name))
    return output
